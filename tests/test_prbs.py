"""Unit and regression tests for per-sample PRBS bit expansion stage.

Governing Specification:
    PA.pdf Section 3.2.4, Equations (3.24)–(3.26), pages 28–29.
    PA.pdf Section 3.2.8.c, Equation (3.64), page 35 (Key Agreement Rate).
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest

from rssi_skg.cli import main, run_correlation_analysis
from rssi_skg.prbs import (
    DEFAULT_POLY_2BIT_MASK,
    DEFAULT_POLY_4BIT_MASK,
    EXCEL_MAX_CELL_CHARS,
    GaloisLFSR,
    PRBSConfig,
    SamplePRBSResult,
    _chunk_bitstream,
    expand_channel_prbs,
    expand_channels_prbs,
    expand_seed,
    export_prbs_bitstreams_csv,
    export_prbs_config_json,
    export_prbs_metrics_csv,
    export_prbs_samples_csv,
    export_prbs_workbook_xlsx,
)
from rssi_skg.quantization import (
    ChannelQuantizationResult,
    QuantizationConfig,
    quantize_channels,
)


# ==============================================================================
# 1. LFSR Primitive Polynomial & Period Verification
# ==============================================================================


def test_prbs_polynomial_verification() -> None:
    """Verify that selected polynomials are primitive and produce maximal length 2^m - 1."""
    # Degree 2: f(x) = x^2 + x + 1, feedback mask 0b11 (3), maximal period = 3
    lfsr2 = GaloisLFSR(degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK)
    assert lfsr2.is_primitive() is True
    for seed in ("01", "10", "11", 1, 2, 3):
        assert lfsr2.calculate_period(seed) == 3
        lfsr2.reset(seed)
        bits = lfsr2.generate_bits(3)
        assert len(bits) == 3

    # Degree 4: f(x) = x^4 + x + 1, feedback mask 0b0011 (3), maximal period = 15
    lfsr4 = GaloisLFSR(degree=4, feedback_mask=DEFAULT_POLY_4BIT_MASK)
    assert lfsr4.is_primitive() is True
    for seed in range(1, 16):
        assert lfsr4.calculate_period(seed) == 15
        lfsr4.reset(seed)
        bits = lfsr4.generate_bits(15)
        assert len(bits) == 15


def test_prbs_zero_state_lockup() -> None:
    """Verify zero-seed handling: lockup at state 0 with period 1 producing all zeros."""
    # Degree 2: 00 -> 000, actual period 1, zero state
    lfsr2 = GaloisLFSR(degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK)
    assert lfsr2.calculate_period(0) == 1
    assert lfsr2.calculate_period("00") == 1
    lfsr2.reset("00")
    bits2 = lfsr2.generate_bits(3)
    assert bits2 == "000"

    # Degree 4: 0000 -> 000000000000000, actual period 1, zero state
    lfsr4 = GaloisLFSR(degree=4, feedback_mask=DEFAULT_POLY_4BIT_MASK)
    assert lfsr4.calculate_period(0) == 1
    assert lfsr4.calculate_period("0000") == 1
    lfsr4.reset("0000")
    bits4 = lfsr4.generate_bits(15)
    assert bits4 == "0" * 15


# ==============================================================================
# 2. Independent Known-Answer Sequences
# ==============================================================================


def test_prbs_known_answer_sequences_2bit() -> None:
    """Verify exact hand-calculated 2-bit Galois LFSR sequences (m=2, mask=3).

    Hand trace for degree 2, mask 0b11:
    - Seed '00' (state 0):
        step 1: out 0, state 0
        step 2: out 0, state 0
        step 3: out 0, state 0 -> '000', period 1
    - Seed '01' (state 1):
        step 1: out 0, next state: (1 << 1) & 3 = 2
        step 2: out 1, next state: ((2 << 1) & 3) ^ 3 = 0 ^ 3 = 3
        step 3: out 1, next state: ((3 << 1) & 3) ^ 3 = 2 ^ 3 = 1 -> '011', period 3
    - Seed '11' (state 3):
        step 1: out 1, next state: ((3 << 1) & 3) ^ 3 = 2 ^ 3 = 1
        step 2: out 0, next state: (1 << 1) & 3 = 2
        step 3: out 1, next state: ((2 << 1) & 3) ^ 3 = 3 -> '101', period 3
    - Seed '10' (state 2):
        step 1: out 1, next state: ((2 << 1) & 3) ^ 3 = 3
        step 2: out 1, next state: ((3 << 1) & 3) ^ 3 = 1
        step 3: out 0, next state: (1 << 1) & 3 = 2 -> '110', period 3
    """
    res_00_bits, res_00_zero, res_00_period, res_00_len = expand_seed("00", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)
    assert res_00_bits == "000"
    assert res_00_period == 1
    assert res_00_zero is True
    assert res_00_len == 3

    res_01_bits, res_01_zero, res_01_period, res_01_len = expand_seed("01", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)
    assert res_01_bits == "011"
    assert res_01_period == 3
    assert res_01_zero is False

    res_11_bits, res_11_zero, res_11_period, res_11_len = expand_seed("11", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)
    assert res_11_bits == "101"
    assert res_11_period == 3
    assert res_11_zero is False

    res_10_bits, res_10_zero, res_10_period, res_10_len = expand_seed("10", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)
    assert res_10_bits == "110"
    assert res_10_period == 3
    assert res_10_zero is False


def test_prbs_known_answer_sequences_4bit() -> None:
    """Verify independent known-answer sequences for all 16 four-bit states (m=4, mask=0b0011).

    Hand-derived state transitions under Galois recurrence S_{k+1} = ((S_k << 1) & 15) ^ (3 if MSB else 0):
    Cycle: 1 -> 2 -> 4 -> 8 -> 3 -> 6 -> 12 -> 11 -> 5 -> 10 -> 7 -> 14 -> 15 -> 13 -> 9 -> 1.
    All non-zero states must produce maximal period 15 with zero_state=False.
    State 0 ('0000') locks up at 0 with actual period 1 and zero_state=True.
    """
    independent_known_answers_4bit = {
        "0000": ("000000000000000", True, 1),
        "0001": ("000100110101111", False, 15),
        "0010": ("001001101011110", False, 15),
        "0011": ("001101011110001", False, 15),
        "0100": ("010011010111100", False, 15),
        "0101": ("010111100010011", False, 15),
        "0110": ("011010111100010", False, 15),
        "0111": ("011110001001101", False, 15),
        "1000": ("100110101111000", False, 15),
        "1001": ("100010011010111", False, 15),
        "1010": ("101111000100110", False, 15),
        "1011": ("101011110001001", False, 15),
        "1100": ("110101111000100", False, 15),
        "1101": ("110001001101011", False, 15),
        "1110": ("111100010011010", False, 15),
        "1111": ("111000100110101", False, 15),
    }

    assert len(independent_known_answers_4bit) == 16

    for seed_str, (expected_bits, expected_zero, expected_period) in independent_known_answers_4bit.items():
        bits, is_zero, actual_period, req_len = expand_seed(
            seed_str,
            degree=4,
            feedback_mask=DEFAULT_POLY_4BIT_MASK,
            length=15,
        )
        assert bits == expected_bits, f"Mismatch for seed {seed_str}: got {bits}, expected {expected_bits}"
        assert is_zero is expected_zero, f"Zero state mismatch for seed {seed_str}"
        assert actual_period == expected_period, f"Period mismatch for seed {seed_str}"
        assert req_len == 15


def test_prbs_requested_length_distinction() -> None:
    """Verify distinction between requested output length and actual period."""
    # Requested length 10 from seed '01' (period 3): should repeat '011' periodically
    bits, is_zero, actual_period, req_len = expand_seed("01", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=10)
    assert bits == "0110110110"
    assert actual_period == 3
    assert req_len == 10

    # Requested length 7 from zero-seed '00' (period 1): should produce '0000000'
    z_bits, z_zero, z_period, z_len = expand_seed("00", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=7)
    assert z_bits == "0000000"
    assert z_period == 1
    assert z_len == 7


# ==============================================================================
# 3. Input Validation and Exception Handling
# ==============================================================================


def test_prbs_input_validation() -> None:
    """Verify strict validation of degree, seed format, and lengths."""
    # Seed length mismatch
    with pytest.raises(ValueError, match="does not match register degree"):
        expand_seed("010", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)

    # Invalid characters in seed
    with pytest.raises(ValueError, match="must contain only '0' and '1'"):
        expand_seed("0a", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)
    with pytest.raises(ValueError, match="must contain only '0' and '1'"):
        expand_seed("12", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)

    # Length < 0
    with pytest.raises(ValueError, match="length must be non-negative"):
        expand_seed("01", degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=-1)

    # Invalid LFSR state
    lfsr = GaloisLFSR(degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK)
    with pytest.raises(ValueError, match="out of range"):
        lfsr.reset(-1)

    # Invalid LFSR mask
    with pytest.raises(ValueError, match="feedback_mask must be within"):
        GaloisLFSR(degree=2, feedback_mask=0)


# ==============================================================================
# 4. Deterministic Per-Sample Reset and Channel Isolation
# ==============================================================================


def test_prbs_per_sample_deterministic_reset() -> None:
    """Verify that LFSR state resets independently for every sample."""
    seeds = ["01", "11", "00", "01"]
    results = [expand_seed(s, degree=2, feedback_mask=DEFAULT_POLY_2BIT_MASK, length=3)[0] for s in seeds]

    # Sample 0 and Sample 3 both had seed '01' -> must have identical outputs
    assert results[0] == results[3] == "011"
    assert results[1] == "101"
    assert results[2] == "000"


def test_prbs_channel_isolation_and_sample_ordering() -> None:
    """Verify channel isolation and strict sample indexing preservation."""
    # Construct mock channel quantization results
    ch_res = expand_channel_prbs(
        levels=[0, 1, 2],
        config=PRBSConfig(),
        channel_name="Alice",
    )

    assert ch_res.sample_count == 3
    assert ch_res.zero_seed_count_2bit == 1
    assert ch_res.zero_seed_count_4bit == 0
    assert ch_res.bitstream_2bit == "000011101"
    assert len(ch_res.bitstream_4bit) == 45

    # Check sample indices
    assert [s.sample_index for s in ch_res.samples] == [0, 1, 2]
    assert ch_res.samples[0].seed_2bit == "00"
    assert ch_res.samples[0].prbs_2bit == "000"
    assert ch_res.samples[0].is_zero_state_2bit is True


# ==============================================================================
# 5. Full Multi-Channel Expansion & KAR Comparison
# ==============================================================================


def test_prbs_multi_channel_expansion_and_kar() -> None:
    """Verify multi-channel expansion and KAR before/after calculation."""
    # Synthesize small 4-sample dataset for Alice, Bob, Eve1-Alice, Eve1-Bob
    series_data = {
        "Alice": [10.0, 20.0, 30.0, 10.0],
        "Bob": [10.0, 20.0, 25.0, 10.0],
        "Eve1-Alice": [20.0, 20.0, 35.0, 35.0],
        "Eve1-Bob": [20.0, 10.0, 20.0, 35.0],
    }
    mock_qres = quantize_channels(
        channels=series_data,
        config=QuantizationConfig(window_size=4, alpha=0.5),
    )

    prbs_res = expand_channels_prbs(quant_result=mock_qres)
    assert len(prbs_res.channels) == 4
    assert len(prbs_res.pair_metrics) == 3

    ab_metric = prbs_res.pair_metrics[0]
    assert ab_metric["pair"] == "Alice–Bob"
    assert ab_metric["bit_count_2bit_before"] == 8
    assert ab_metric["bit_count_2bit_after"] == 12  # 4 * 3
    assert ab_metric["bit_count_4bit_before"] == 16
    assert ab_metric["bit_count_4bit_after"] == 60  # 4 * 15


# ==============================================================================
# 6. Artifact Exports Validation
# ==============================================================================


def test_prbs_artifact_exports(tmp_path: Path) -> None:
    """Test exporting all 5 PRBS artifacts: samples, bitstreams, metrics, JSON, and XLSX."""
    series_data = {
        "Alice": [10.0, 25.0],
        "Bob": [10.0, 25.0],
        "Eve1-Alice": [25.0, 25.0],
        "Eve1-Bob": [25.0, 25.0],
    }
    qres = quantize_channels(
        channels=series_data,
        config=QuantizationConfig(window_size=2, alpha=0.5),
    )

    prbs_res = expand_channels_prbs(quant_result=qres)

    # 1. Samples CSV
    csv_samples = tmp_path / "prbs_samples.csv"
    export_prbs_samples_csv(prbs_res, csv_samples)
    assert csv_samples.exists()
    with open(csv_samples, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        assert len(rows) == 2  # 2 samples
        # Verify columns exist and format
        assert rows[0]["Alice_seed_2bit"] in ("00", "01", "11")
        assert len(rows[0]["Alice_prbs_2bit"]) == 3

    # 2. Bitstreams CSV
    csv_bitstreams = tmp_path / "prbs_bitstreams.csv"
    export_prbs_bitstreams_csv(prbs_res, csv_bitstreams)
    assert csv_bitstreams.exists()
    with open(csv_bitstreams, "r", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
        assert len(rows) == 8  # 4 channels * 2 encoding paths
        assert rows[0]["channel"] == "Alice"
        assert len(rows[0]["bitstream"]) == 6  # 2 samples * 3 bits

    # 3. Metrics CSV
    csv_metrics = tmp_path / "prbs_metrics.csv"
    export_prbs_metrics_csv(prbs_res, csv_metrics)
    assert csv_metrics.exists()
    with open(csv_metrics, "r", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
        assert len(rows) == 3
        assert "kar_2bit_after" in rows[0]

    # 4. Config JSON
    json_path = tmp_path / "prbs_config.json"
    export_prbs_config_json(prbs_res, json_path, extra_metadata={"note": "test"})
    assert json_path.exists()
    with open(json_path, "r", encoding="utf-8") as f:
        meta = json.load(f)
        assert meta["note"] == "test"
        assert meta["implementation_choices"]["register_sizes"]["2bit_path"]["degree_m"] == 2
        assert "zero_seed_handling" in meta["implementation_choices"]

    # 5. Workbook XLSX
    xlsx_path = tmp_path / "prbs.xlsx"
    export_prbs_workbook_xlsx(prbs_res, xlsx_path)
    assert xlsx_path.exists()
    wb = openpyxl.load_workbook(xlsx_path)
    assert "PRBS Samples" in wb.sheetnames
    assert "PRBS Bitstreams" in wb.sheetnames
    assert "KAR Comparison" in wb.sheetnames
    assert "Configuration & Zero Seeds" in wb.sheetnames


def test_excel_cell_chunking() -> None:
    """Verify bitstream chunking for strings exceeding Excel's 32,767 character limit."""
    short_bitstream = "101" * 10
    assert _chunk_bitstream(short_bitstream) == [short_bitstream]

    long_bitstream = "1" * (EXCEL_MAX_CELL_CHARS + 50)
    chunks = _chunk_bitstream(long_bitstream)
    assert len(chunks) == 2
    assert len(chunks[0]) == EXCEL_MAX_CELL_CHARS
    assert len(chunks[1]) == 50


# ==============================================================================
# 7. CLI End-to-End Pipeline Integration on Dummy Dataset
# ==============================================================================


def test_cli_prbs_pipeline_integration(tmp_path: Path) -> None:
    """Run full CLI pipeline with --prbs on dummy dataset and verify metrics."""
    input_xlsx = Path("data/raw/dummy/Dummy RSSI (data Indoor).xlsx")
    assert input_xlsx.exists(), "Dummy workbook missing"

    out_csv = tmp_path / "correlation.csv"
    samples_csv = tmp_path / "prbs_samples.csv"
    bitstreams_csv = tmp_path / "prbs_bitstreams.csv"
    metrics_csv = tmp_path / "prbs_metrics.csv"
    config_json = tmp_path / "prbs_config.json"
    prbs_xlsx = tmp_path / "prbs.xlsx"

    argv = [
        "-i",
        str(input_xlsx),
        "-o",
        str(out_csv),
        "--prbs",
        "--prbs-samples-output",
        str(samples_csv),
        "--prbs-bitstreams-output",
        str(bitstreams_csv),
        "--prbs-metrics-output",
        str(metrics_csv),
        "--prbs-config-output",
        str(config_json),
        "--prbs-xlsx-output",
        str(prbs_xlsx),
    ]

    ret = main(argv)
    assert ret == 0

    assert samples_csv.exists()
    assert bitstreams_csv.exists()
    assert metrics_csv.exists()
    assert config_json.exists()
    assert prbs_xlsx.exists()

    # Verify bitstream lengths for 500 samples
    # 2-bit: 500 * 3 = 1,500 bits
    # 4-bit: 500 * 15 = 7,500 bits
    with open(bitstreams_csv, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        b_rows = list(reader)
        assert len(b_rows) == 8
        for row in b_rows:
            assert int(row["sample_count"]) == 500
            if "2-bit" in row["encoding_path"]:
                assert len(row["bitstream"]) == 1500
            elif "4-bit" in row["encoding_path"]:
                assert len(row["bitstream"]) == 7500

    # Verify metrics matching known values
    with open(metrics_csv, "r", encoding="utf-8") as f:
        m_rows = list(csv.DictReader(f))
        assert len(m_rows) == 3
        # Check Alice-Bob KAR
        ab = m_rows[0]
        assert ab["pair"] == "Alice–Bob"
        assert float(ab["kar_2bit_before"]) == pytest.approx(80.4000, abs=1e-3)
        assert float(ab["kar_2bit_after"]) == pytest.approx(76.2667, abs=1e-3)
        assert float(ab["kar_4bit_before"]) == pytest.approx(90.2000, abs=1e-3)
        assert float(ab["kar_4bit_after"]) == pytest.approx(81.0133, abs=1e-3)
