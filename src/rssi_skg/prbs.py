"""Pseudo-Random Binary Sequence (PRBS) bit expansion module using Galois XOR-LFSR.

Governing Specification:
    PA.pdf Section 3.2.4, Equations (3.24)–(3.26), pages 28–29.
    PA.pdf Section 3.2.8.c, Equation (3.64), page 35 (Key Agreement Rate).

Mathematical Formulation:
    Polynomial representation (Equation 3.24):
        f(x) = g_n x^n + g_{n-1} x^{n-1} + ... + g_0 x^0
        where n is the register length (degree m) and g_i in {0, 1} are feedback coefficients.

    Bit recurrence relation (Equation 3.25):
        S_{i+m} = (sum_{j=0}^{m-1} C_j S_{i+j}) mod 2

    Maximal output period for non-zero seeds with primitive polynomial (Equation 3.26):
        L = 2^m - 1

Implementation Choices & Design Decisions:
    1. Register Sizes:
       - 2-bit path: degree m = 2, generates exactly 2^2 - 1 = 3 bits per sample.
       - 4-bit path: degree m = 4, generates exactly 2^4 - 1 = 15 bits per sample.
       Both register sizes and their 1-to-1 correspondence with quantization codewords
       are implementation choices chosen to align with the 2-bit Gray and 4-bit representations.
    2. Primitive Polynomials & Galois Feedback Masks:
       - Degree 2: f(x) = x^2 + x + 1 (unique primitive polynomial of degree 2 over GF(2)).
         In Left-shift Galois LFSR, feedback is taken from MSB (r_1) and XORed into lower taps.
         Mask = 0b11 (decimal 3), corresponding to x^1 + x^0.
       - Degree 4: f(x) = x^4 + x + 1 (standard primitive polynomial of degree 4 over GF(2)).
         In Left-shift Galois LFSR, feedback is taken from MSB (r_3) and XORed into lower taps.
         Mask = 0b0011 (decimal 3), corresponding to x^1 + x^0.
    3. Shift Direction & Bit Ordering:
       - Shift Direction: Left-shift (bits shift toward MSB; MSB is emitted as the output bit).
       - Bit Ordering: Big-endian representation where index 0 of the binary seed string is MSB.
    4. Output Timing:
       - The current MSB state bit is emitted as the output bit for step t, then the register
         executes the internal Galois XOR shift to transition to state t+1.
    5. Per-Sample Resets:
       - For every individual sample, the shift register is deterministically reset to that sample's
         seed codeword. The LFSR runs for exactly 2^m - 1 steps, emitting 1 bit per step.
       - The resulting bit segments are concatenated strictly in original sample order.
    6. Zero-Seed Handling (Deliberate Exception to PA.pdf's Non-Zero Seed Condition):
       - PA.pdf Equation (3.26) states that maximal-period sequence generation requires a non-zero seed.
       - In this research experiment, the 2-bit Gray encoding maps quantization level 0 to "00".
       - Per explicit experimental design, seed "00" is NEVER remapped to "10" or any other value.
       - Standard Galois XOR-LFSR recurrence is executed unmodified: state 0 shifts to 0, producing
         exactly "000" for m=2 (and "0"*15 if an all-zero seed were supplied for m=4).
       - Zero-state lockup is formally recorded: is_zero_state = True, actual_period = 1,
         distinguished from requested_length = 2^m - 1.
    7. Scientific Limitation & Entropy Caveat:
       - PRBS bit expansion is a deterministic linear state machine.
       - Expanding bits does NOT increase physical source entropy or cryptographic security.
       - Repeated quantization codewords produce identical expanded sequences.
       - Eavesdropper or legitimate KAR may decrease following expansion due to maximal dispersion
         of differing codewords across 2^m - 1 bit positions.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
import json
import math
from pathlib import Path
from typing import Any, Mapping, Sequence

import openpyxl

from rssi_skg.quantization import (
    CODE_4BIT_MAP,
    EVALUATION_PAIRS,
    EXCEL_MAX_CELL_CHARS,
    GRAY_2BIT_MAP,
    REQUIRED_CHANNELS,
    ChannelQuantizationResult,
    calculate_kar,
    calculate_level_agreement,
)

# Default primitive polynomials over GF(2)
# Degree 2: f(x) = x^2 + x + 1 -> taps x^1 + x^0 -> mask 0b11 (3)
DEFAULT_POLY_2BIT_NAME: str = "x^2 + x + 1"
DEFAULT_POLY_2BIT_DEGREE: int = 2
DEFAULT_POLY_2BIT_MASK: int = 0b11

# Degree 4: f(x) = x^4 + x + 1 -> taps x^1 + x^0 -> mask 0b0011 (3)
DEFAULT_POLY_4BIT_NAME: str = "x^4 + x + 1"
DEFAULT_POLY_4BIT_DEGREE: int = 4
DEFAULT_POLY_4BIT_MASK: int = 0b0011


class GaloisLFSR:
    """Galois XOR Linear Feedback Shift Register (LFSR) implementation.

    Executes Galois (internal XOR) state transitions over GF(2).
    """

    def __init__(
        self,
        degree: int,
        feedback_mask: int,
        shift_direction: str = "left",
    ) -> None:
        """Initialize Galois LFSR.

        Args:
            degree: Register length m (number of flip-flops/bits).
            feedback_mask: Bitmask representing internal XOR feedback taps (excluding x^m).
            shift_direction: Direction of shift ("left" or "right"). Default is "left".

        Raises:
            TypeError: If degree or feedback_mask is not an integer.
            ValueError: If degree < 1, feedback_mask is outside [1, (1<<degree)-1],
                or shift_direction is not "left" or "right".
        """
        if isinstance(degree, bool) or not isinstance(degree, int):
            raise TypeError(f"degree must be an integer, got {type(degree).__name__}.")
        if degree < 1:
            raise ValueError(f"degree must be at least 1, got {degree}.")

        if isinstance(feedback_mask, bool) or not isinstance(feedback_mask, int):
            raise TypeError(
                f"feedback_mask must be an integer, got {type(feedback_mask).__name__}."
            )
        max_mask = (1 << degree) - 1
        if not (1 <= feedback_mask <= max_mask):
            raise ValueError(
                f"feedback_mask must be within [1, {max_mask}], got {feedback_mask}."
            )

        if shift_direction not in ("left", "right"):
            raise ValueError(
                f"shift_direction must be 'left' or 'right', got {shift_direction!r}."
            )

        self._degree = degree
        self._feedback_mask = feedback_mask
        self._shift_direction = shift_direction
        self._full_mask = max_mask
        self._state: int = 0
        self._initial_state: int = 0

    @property
    def degree(self) -> int:
        """Register degree (m)."""
        return self._degree

    @property
    def feedback_mask(self) -> int:
        """Galois XOR feedback mask."""
        return self._feedback_mask

    @property
    def shift_direction(self) -> str:
        """Shift direction ('left' or 'right')."""
        return self._shift_direction

    @property
    def state(self) -> int:
        """Current internal register state."""
        return self._state

    def reset(self, seed: int | str) -> None:
        """Reset internal register state to the given seed.

        Args:
            seed: Binary string (e.g. '01', '1010') or non-negative integer.

        Raises:
            TypeError: If seed is not int or str.
            ValueError: If string seed length does not match degree, contains non-binary
                characters, or integer seed is outside [0, (1<<degree)-1].
        """
        if isinstance(seed, str):
            if len(seed) != self._degree:
                raise ValueError(
                    f"Seed string length {len(seed)} does not match register degree {self._degree}."
                )
            if not all(c in "01" for c in seed):
                raise ValueError(f"Seed string must contain only '0' and '1', got {seed!r}.")
            int_seed = int(seed, 2)
        elif isinstance(seed, int) and not isinstance(seed, bool):
            if not (0 <= seed <= self._full_mask):
                raise ValueError(
                    f"Integer seed {seed} is out of range [0, {self._full_mask}] for degree {self._degree}."
                )
            int_seed = seed
        else:
            raise TypeError(f"Seed must be an int or str, got {type(seed).__name__}.")

        self._state = int_seed
        self._initial_state = int_seed

    def step(self) -> int:
        """Execute one clock cycle and return the emitted output bit.

        In left-shift:
            - Output bit is MSB: (state >> (degree - 1)) & 1.
            - Next state: (state << 1) & full_mask.
            - If output bit was 1, XOR next state with feedback_mask.

        In right-shift:
            - Output bit is LSB: state & 1.
            - Next state: state >> 1.
            - If output bit was 1, XOR next state with feedback_mask.

        Returns:
            int: Emitted output bit (0 or 1).
        """
        if self._shift_direction == "left":
            output_bit = (self._state >> (self._degree - 1)) & 1
            next_state = (self._state << 1) & self._full_mask
            if output_bit == 1:
                next_state ^= self._feedback_mask
            self._state = next_state
            return output_bit
        else:
            output_bit = self._state & 1
            next_state = self._state >> 1
            if output_bit == 1:
                next_state ^= self._feedback_mask
            self._state = next_state
            return output_bit

    def generate_bits(self, count: int) -> str:
        """Generate a sequence of output bits from current state.

        Args:
            count: Number of output bits to generate (must be >= 0).

        Returns:
            str: Binary bitstream string of length `count`.

        Raises:
            ValueError: If count < 0.
        """
        if count < 0:
            raise ValueError(f"Bit count must be non-negative, got {count}.")
        bits = [str(self.step()) for _ in range(count)]
        return "".join(bits)

    def calculate_period(self, seed: int | str) -> int:
        """Calculate the actual state cycle period from a given seed.

        For a zero seed (state 0), standard evolution remains in state 0, so period = 1.
        For non-zero seeds with a primitive polynomial, period = (1 << degree) - 1.

        Args:
            seed: Seed value (int or str).

        Returns:
            int: Actual cycle period in steps.
        """
        self.reset(seed)
        start_state = self._state
        if start_state == 0:
            return 1

        max_steps = 1 << self._degree
        for step_idx in range(1, max_steps + 2):
            self.step()
            if self._state == start_state:
                return step_idx

        return -1

    def verify_primitive(self) -> bool:
        """Verify whether the feedback polynomial yields maximal period (2^m - 1) for all non-zero seeds.

        Returns:
            bool: True if maximal period (2^m - 1) is achieved for all non-zero states.
        """
        maximal_period = (1 << self._degree) - 1
        for s in range(1, 1 << self._degree):
            if self.calculate_period(s) != maximal_period:
                return False
        return True

    def is_primitive(self) -> bool:
        """Alias for verify_primitive()."""
        return self.verify_primitive()


def expand_seed(
    seed: str,
    degree: int,
    feedback_mask: int,
    shift_direction: str = "left",
    length: int | None = None,
) -> tuple[str, bool, int, int]:
    """Expand a single binary seed string using Galois XOR-LFSR.

    Args:
        seed: Binary seed string of length `degree` (e.g. '00', '01', '1010').
        degree: Register degree m.
        feedback_mask: Galois feedback mask.
        shift_direction: 'left' (default) or 'right'.
        length: Requested number of output bits. If None, defaults to 2^degree - 1.

    Returns:
        tuple[str, bool, int, int]:
            - expanded_bits: Emitted binary string.
            - is_zero_state: True if seed is all zeros (zero-state lockup).
            - actual_period: Actual cycle period (1 for zero seed, 2^m - 1 for non-zero primitive).
            - requested_length: Requested output bit count.

    Raises:
        ValueError: If seed length does not match degree or contains non-binary characters.
    """
    if length is None:
        requested_len = (1 << degree) - 1
    else:
        if length < 0:
            raise ValueError(f"Requested length must be non-negative, got {length}.")
        requested_len = length

    lfsr = GaloisLFSR(
        degree=degree,
        feedback_mask=feedback_mask,
        shift_direction=shift_direction,
    )
    lfsr.reset(seed)

    int_seed = int(seed, 2)
    is_zero = (int_seed == 0)
    actual_period = 1 if is_zero else ((1 << degree) - 1)
    expanded_bits = lfsr.generate_bits(requested_len)

    return expanded_bits, is_zero, actual_period, requested_len


@dataclass(frozen=True)
class PRBSConfig:
    """Configuration parameters for per-sample PRBS bit expansion.

    Attributes:
        poly_2bit_name: Mathematical formula for 2-bit polynomial.
        degree_2bit: Degree m for 2-bit path (default: 2).
        mask_2bit: Galois feedback mask for 2-bit path (default: 0b11 = 3).
        poly_4bit_name: Mathematical formula for 4-bit polynomial.
        degree_4bit: Degree m for 4-bit path (default: 4).
        mask_4bit: Galois feedback mask for 4-bit path (default: 0b0011 = 3).
        shift_direction: Shift direction ("left" or "right"). Default: "left".
    """

    poly_2bit_name: str = DEFAULT_POLY_2BIT_NAME
    degree_2bit: int = DEFAULT_POLY_2BIT_DEGREE
    mask_2bit: int = DEFAULT_POLY_2BIT_MASK
    poly_4bit_name: str = DEFAULT_POLY_4BIT_NAME
    degree_4bit: int = DEFAULT_POLY_4BIT_DEGREE
    mask_4bit: int = DEFAULT_POLY_4BIT_MASK
    shift_direction: str = "left"

    def __post_init__(self) -> None:
        if self.degree_2bit != 2:
            raise ValueError(f"degree_2bit must be 2, got {self.degree_2bit}.")
        if self.degree_4bit != 4:
            raise ValueError(f"degree_4bit must be 4, got {self.degree_4bit}.")
        if self.shift_direction not in ("left", "right"):
            raise ValueError(
                f"shift_direction must be 'left' or 'right', got {self.shift_direction!r}."
            )

    @property
    def output_length_2bit(self) -> int:
        """Requested output bits per sample for 2-bit path (2^2 - 1 = 3)."""
        return (1 << self.degree_2bit) - 1

    @property
    def output_length_4bit(self) -> int:
        """Requested output bits per sample for 4-bit path (2^4 - 1 = 15)."""
        return (1 << self.degree_4bit) - 1

    def to_metadata_dict(self) -> dict[str, Any]:
        """Convert configuration to structured metadata dictionary."""
        return {
            "algorithm": "Per-Sample Galois XOR-LFSR PRBS Expansion",
            "governing_specification": {
                "lfsr_expansion": "PA.pdf Section 3.2.4, Equations (3.24)–(3.26), pages 28–29",
                "key_agreement_rate": "PA.pdf Section 3.2.8.c, Equation (3.64), page 35",
            },
            "implementation_choices": {
                "register_sizes": {
                    "2bit_path": {
                        "degree_m": self.degree_2bit,
                        "polynomial": self.poly_2bit_name,
                        "galois_mask": f"0b{self.mask_2bit:0{self.degree_2bit}b}",
                        "galois_mask_dec": self.mask_2bit,
                        "bits_per_sample": self.output_length_2bit,
                        "status": "primitive polynomial over GF(2) with maximal period 2^2 - 1 = 3",
                    },
                    "4bit_path": {
                        "degree_m": self.degree_4bit,
                        "polynomial": self.poly_4bit_name,
                        "galois_mask": f"0b{self.mask_4bit:0{self.degree_4bit}b}",
                        "galois_mask_dec": self.mask_4bit,
                        "bits_per_sample": self.output_length_4bit,
                        "status": "primitive polynomial over GF(2) with maximal period 2^4 - 1 = 15",
                    },
                },
                "shift_direction": self.shift_direction,
                "bit_ordering": "big-endian (index 0 of seed string is MSB)",
                "output_timing": "MSB emitted at current clock step before internal Galois XOR transition",
                "per_sample_resets": True,
                "zero_seed_handling": (
                    "Deliberate exception to PA.pdf's non-zero seed condition: seed '00' is NEVER remapped; "
                    "standard Galois XOR-LFSR recurrence runs unmodified, producing '000' with actual period 1."
                ),
            },
            "scientific_limitations_and_security_notice": {
                "entropy_warning": (
                    "PRBS bit expansion is a deterministic linear state machine. "
                    "Longer output length does NOT add physical channel entropy or cryptographic security. "
                    "Repeated quantization codewords produce identical PRBS sequences, and KAR may decrease."
                ),
            },
        }


@dataclass(frozen=True)
class SamplePRBSResult:
    """Expansion result for a single sample across both 2-bit and 4-bit representations."""

    sample_index: int
    level: int
    encoding_2bit: str
    seed_2bit: str
    prbs_2bit: str
    is_zero_state_2bit: bool
    actual_period_2bit: int
    requested_length_2bit: int
    encoding_4bit: str
    seed_4bit: str
    prbs_4bit: str
    is_zero_state_4bit: bool
    actual_period_4bit: int
    requested_length_4bit: int


@dataclass(frozen=True)
class ChannelPRBSResult:
    """PRBS expansion result for an individual communication channel."""

    channel_name: str
    sample_count: int
    samples: list[SamplePRBSResult]
    bitstream_2bit: str
    bitstream_4bit: str
    zero_seed_count_2bit: int
    zero_seed_count_4bit: int

    @property
    def zero_seed_percentage_2bit(self) -> float:
        """Percentage of samples with 2-bit zero seed '00'."""
        return (self.zero_seed_count_2bit / self.sample_count * 100.0) if self.sample_count > 0 else 0.0

    @property
    def zero_seed_percentage_4bit(self) -> float:
        """Percentage of samples with 4-bit zero seed '0000'."""
        return (self.zero_seed_count_4bit / self.sample_count * 100.0) if self.sample_count > 0 else 0.0


def _chunk_bitstream(stream: str, chunk_size: int = EXCEL_MAX_CELL_CHARS) -> list[str]:
    """Split a bitstream string into chunks of maximum length chunk_size.

    Args:
        stream: String bitstream.
        chunk_size: Maximum characters per chunk (default: 32,767 for Excel).

    Returns:
        list[str]: List of string chunks.
    """
    if len(stream) <= chunk_size:
        return [stream]
    return [stream[i : i + chunk_size] for i in range(0, len(stream), chunk_size)]


@dataclass
class MultiChannelPRBSResult:
    """Complete multi-channel PRBS expansion result with before/after KAR metrics."""

    config: PRBSConfig
    sample_count: int
    channels: dict[str, ChannelPRBSResult]
    pair_metrics: list[dict[str, Any]]
    figure_paths: tuple[Path, Path] | None = None


def expand_channel_prbs(
    levels: Sequence[int],
    config: PRBSConfig,
    channel_name: str = "Alice",
) -> ChannelPRBSResult:
    """Expand quantization levels of a single channel into 2-bit and 4-bit PRBS bitstreams.

    Args:
        levels: Sequence of discrete quantization levels (0, 1, 2).
        config: PRBS configuration.
        channel_name: Name of the channel.

    Returns:
        ChannelPRBSResult: Fully populated channel PRBS expansion container.
    """
    sample_results: list[SamplePRBSResult] = []
    bitstream_2bit_chunks: list[str] = []
    bitstream_4bit_chunks: list[str] = []
    zero_seed_count_2b = 0
    zero_seed_count_4b = 0

    len_2b = config.output_length_2bit
    len_4b = config.output_length_4bit

    for idx, lvl in enumerate(levels):
        if lvl not in GRAY_2BIT_MAP or lvl not in CODE_4BIT_MAP:
            raise ValueError(
                f"Invalid quantization level {lvl!r} at sample {idx} in channel '{channel_name}'."
            )

        enc_2b = GRAY_2BIT_MAP[lvl]
        enc_4b = CODE_4BIT_MAP[lvl]

        # Expand 2-bit path
        bits_2b, is_zero_2b, period_2b, req_len_2b = expand_seed(
            seed=enc_2b,
            degree=config.degree_2bit,
            feedback_mask=config.mask_2bit,
            shift_direction=config.shift_direction,
            length=len_2b,
        )
        if is_zero_2b:
            zero_seed_count_2b += 1

        # Expand 4-bit path
        bits_4b, is_zero_4b, period_4b, req_len_4b = expand_seed(
            seed=enc_4b,
            degree=config.degree_4bit,
            feedback_mask=config.mask_4bit,
            shift_direction=config.shift_direction,
            length=len_4b,
        )
        if is_zero_4b:
            zero_seed_count_4b += 1

        sample_res = SamplePRBSResult(
            sample_index=idx,
            level=lvl,
            encoding_2bit=enc_2b,
            seed_2bit=enc_2b,
            prbs_2bit=bits_2b,
            is_zero_state_2bit=is_zero_2b,
            actual_period_2bit=period_2b,
            requested_length_2bit=req_len_2b,
            encoding_4bit=enc_4b,
            seed_4bit=enc_4b,
            prbs_4bit=bits_4b,
            is_zero_state_4bit=is_zero_4b,
            actual_period_4bit=period_4b,
            requested_length_4bit=req_len_4b,
        )
        sample_results.append(sample_res)
        bitstream_2bit_chunks.append(bits_2b)
        bitstream_4bit_chunks.append(bits_4b)

    return ChannelPRBSResult(
        channel_name=channel_name,
        sample_count=len(levels),
        samples=sample_results,
        bitstream_2bit="".join(bitstream_2bit_chunks),
        bitstream_4bit="".join(bitstream_4bit_chunks),
        zero_seed_count_2bit=zero_seed_count_2b,
        zero_seed_count_4bit=zero_seed_count_4b,
    )


def expand_channels_prbs(
    quant_result: ChannelQuantizationResult,
    config: PRBSConfig | None = None,
) -> MultiChannelPRBSResult:
    """Expand quantization results across all channels with independent isolation and pair evaluation.

    Args:
        quant_result: Output from quantize_channels.
        config: PRBSConfig instance (uses defaults if None).

    Returns:
        MultiChannelPRBSResult: Container with expanded channels and before/after KAR metrics.
    """
    cfg = config if config is not None else PRBSConfig()
    channel_results: dict[str, ChannelPRBSResult] = {}

    for ch_name, levels in quant_result.levels.items():
        channel_results[ch_name] = expand_channel_prbs(
            levels=levels,
            config=cfg,
            channel_name=ch_name,
        )

    # Evaluate pair metrics before vs after PRBS expansion
    pair_metrics: list[dict[str, Any]] = []
    # Build lookup of quantization metrics
    quant_pair_lookup: dict[str, dict[str, Any]] = {
        pm["pair"]: pm for pm in quant_result.pair_metrics
    }

    for col1, col2, pair_label in EVALUATION_PAIRS:
        if col1 in channel_results and col2 in channel_results:
            q_pm = quant_pair_lookup.get(pair_label, {})
            lvl_agr = q_pm.get("level_agreement_rate", 0.0)

            orig_kar_2b = q_pm.get("kar_2bit", 0.0)
            orig_kar_4b = q_pm.get("kar_4bit", 0.0)

            prbs_b2_a = channel_results[col1].bitstream_2bit
            prbs_b2_b = channel_results[col2].bitstream_2bit
            prbs_kar_2b = calculate_kar(prbs_b2_a, prbs_b2_b)

            prbs_b4_a = channel_results[col1].bitstream_4bit
            prbs_b4_b = channel_results[col2].bitstream_4bit
            prbs_kar_4b = calculate_kar(prbs_b4_a, prbs_b4_b)

            mismatches_2b = sum(1 for a, b in zip(prbs_b2_a, prbs_b2_b) if a != b)
            matches_2b = len(prbs_b2_a) - mismatches_2b

            mismatches_4b = sum(1 for a, b in zip(prbs_b4_a, prbs_b4_b) if a != b)
            matches_4b = len(prbs_b4_a) - mismatches_4b

            pair_metrics.append(
                {
                    "pair": pair_label,
                    "channel_a": col1,
                    "channel_b": col2,
                    "sample_count": quant_result.sample_count,
                    "level_agreement_rate": lvl_agr,
                    # 2-bit metrics
                    "kar_2bit_before": orig_kar_2b,
                    "kar_2bit_after": prbs_kar_2b,
                    "kar_2bit_delta": prbs_kar_2b - orig_kar_2b,
                    "bit_count_2bit_before": len(quant_result.bitstreams_2bit[col1]),
                    "bit_count_2bit_after": len(prbs_b2_a),
                    "pre_prbs_kar_2bit": orig_kar_2b,
                    "post_prbs_kar_2bit": prbs_kar_2b,
                    "kar_delta_2bit": prbs_kar_2b - orig_kar_2b,
                    "pre_prbs_bits_2bit": len(quant_result.bitstreams_2bit[col1]),
                    "post_prbs_bits_2bit": len(prbs_b2_a),
                    "matches_2bit_after": matches_2b,
                    "mismatches_2bit_after": mismatches_2b,
                    # 4-bit metrics
                    "kar_4bit_before": orig_kar_4b,
                    "kar_4bit_after": prbs_kar_4b,
                    "kar_4bit_delta": prbs_kar_4b - orig_kar_4b,
                    "bit_count_4bit_before": len(quant_result.bitstreams_4bit[col1]),
                    "bit_count_4bit_after": len(prbs_b4_a),
                    "pre_prbs_kar_4bit": orig_kar_4b,
                    "post_prbs_kar_4bit": prbs_kar_4b,
                    "kar_delta_4bit": prbs_kar_4b - orig_kar_4b,
                    "pre_prbs_bits_4bit": len(quant_result.bitstreams_4bit[col1]),
                    "post_prbs_bits_4bit": len(prbs_b4_a),
                    "matches_4bit_after": matches_4b,
                    "mismatches_4bit_after": mismatches_4b,
                    # Zero seed counts
                    "zero_seeds_a_2bit": channel_results[col1].zero_seed_count_2bit,
                    "zero_seeds_b_2bit": channel_results[col2].zero_seed_count_2bit,
                    "zero_seeds_a_4bit": channel_results[col1].zero_seed_count_4bit,
                    "zero_seeds_b_4bit": channel_results[col2].zero_seed_count_4bit,
                }
            )

    return MultiChannelPRBSResult(
        config=cfg,
        sample_count=quant_result.sample_count,
        channels=channel_results,
        pair_metrics=pair_metrics,
    )


# ---------------------------------------------------------------------------
# Artifact Export Functions
# ---------------------------------------------------------------------------


def export_prbs_samples_csv(
    prbs_result: MultiChannelPRBSResult,
    output_path: str | Path,
) -> Path:
    """Export aligned sample-level PRBS expansion details to a CSV file.

    Columns:
        sample_index,
        {channel}_level, {channel}_seed_2bit, {channel}_prbs_2bit, {channel}_zero_state_2bit,
        {channel}_seed_4bit, {channel}_prbs_4bit, {channel}_zero_state_4bit
        for each channel in REQUIRED_CHANNELS.

    Args:
        prbs_result: MultiChannelPRBSResult instance.
        output_path: Destination CSV file path.

    Returns:
        Path: Output file path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    channels = [ch for ch in REQUIRED_CHANNELS if ch in prbs_result.channels]
    header = ["sample_index"]
    for ch in channels:
        header.extend(
            [
                f"{ch}_level",
                f"{ch}_seed_2bit",
                f"{ch}_prbs_2bit",
                f"{ch}_zero_state_2bit",
                f"{ch}_seed_4bit",
                f"{ch}_prbs_4bit",
                f"{ch}_zero_state_4bit",
            ]
        )

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for i in range(prbs_result.sample_count):
            row: list[Any] = [i]
            for ch in channels:
                s_res = prbs_result.channels[ch].samples[i]
                row.extend(
                    [
                        s_res.level,
                        s_res.seed_2bit,
                        s_res.prbs_2bit,
                        s_res.is_zero_state_2bit,
                        s_res.seed_4bit,
                        s_res.prbs_4bit,
                        s_res.is_zero_state_4bit,
                    ]
                )
            writer.writerow(row)

    return out_path


def export_prbs_bitstreams_csv(
    prbs_result: MultiChannelPRBSResult,
    output_path: str | Path,
) -> Path:
    """Export concatenated PRBS bitstreams to a CSV file.

    Columns:
        channel, encoding_path, sample_count, bit_length, zero_seed_count, bitstream

    Args:
        prbs_result: MultiChannelPRBSResult instance.
        output_path: Destination CSV file path.

    Returns:
        Path: Output file path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    header = [
        "channel",
        "encoding_path",
        "sample_count",
        "bit_length",
        "zero_seed_count",
        "bitstream",
    ]

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for ch_name, ch_res in prbs_result.channels.items():
            # 2-bit path
            writer.writerow(
                [
                    ch_name,
                    "2-bit Gray PRBS (m=2)",
                    ch_res.sample_count,
                    len(ch_res.bitstream_2bit),
                    ch_res.zero_seed_count_2bit,
                    ch_res.bitstream_2bit,
                ]
            )
            # 4-bit path
            writer.writerow(
                [
                    ch_name,
                    "4-bit PRBS (m=4)",
                    ch_res.sample_count,
                    len(ch_res.bitstream_4bit),
                    ch_res.zero_seed_count_4bit,
                    ch_res.bitstream_4bit,
                ]
            )

    return out_path


def export_prbs_metrics_csv(
    prbs_result: MultiChannelPRBSResult,
    output_path: str | Path,
) -> Path:
    """Export before/after KAR and agreement metrics comparison to a CSV file.

    Args:
        prbs_result: MultiChannelPRBSResult instance.
        output_path: Destination CSV file path.

    Returns:
        Path: Output file path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = [
        "pair",
        "channel_a",
        "channel_b",
        "sample_count",
        "level_agreement_rate",
        "kar_2bit_before",
        "kar_2bit_after",
        "kar_2bit_delta",
        "bit_count_2bit_before",
        "bit_count_2bit_after",
        "kar_4bit_before",
        "kar_4bit_after",
        "kar_4bit_delta",
        "bit_count_4bit_before",
        "bit_count_4bit_after",
    ]

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for pm in prbs_result.pair_metrics:
            csv_row = {k: pm[k] for k in fieldnames}
            writer.writerow(csv_row)

    return out_path


def export_prbs_config_json(
    prbs_result: MultiChannelPRBSResult,
    output_path: str | Path,
    extra_metadata: dict[str, Any] | None = None,
) -> Path:
    """Export PRBS configuration and metadata to a JSON file.

    Args:
        prbs_result: MultiChannelPRBSResult instance.
        output_path: Destination JSON file path.
        extra_metadata: Additional metadata dictionary.

    Returns:
        Path: Output file path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    meta = prbs_result.config.to_metadata_dict()
    meta["sample_count"] = prbs_result.sample_count

    meta["channels_summary"] = {}
    for ch_name, ch_res in prbs_result.channels.items():
        meta["channels_summary"][ch_name] = {
            "sample_count": ch_res.sample_count,
            "zero_seeds_2bit": ch_res.zero_seed_count_2bit,
            "zero_seeds_4bit": ch_res.zero_seed_count_4bit,
            "bitstream_length_2bit": len(ch_res.bitstream_2bit),
            "bitstream_length_4bit": len(ch_res.bitstream_4bit),
        }

    meta["pair_metrics_summary"] = prbs_result.pair_metrics

    if extra_metadata:
        meta.update(extra_metadata)

    with open(out_path, mode="w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    return out_path


def export_prbs_workbook_xlsx(
    prbs_result: MultiChannelPRBSResult,
    output_path: str | Path,
) -> Path:
    """Export companion multi-tab Excel workbook (.xlsx) preserving leading zeros and long bitstrings.

    Sheets:
        1. 'PRBS Samples': Sample-level index, discrete level, seeds, expanded bits, zero-state status.
           Bitstrings explicitly stored as Text (data_type='s', number_format='@').
        2. 'PRBS Bitstreams': Concatenated bitstreams with cell length handling (EXCEL_MAX_CELL_CHARS).
        3. 'KAR Comparison': Pre vs Post PRBS KAR metrics comparison.
        4. 'Configuration & Zero Seeds': Configuration, Galois polynomials, masks, and zero seed counts.

    Args:
        prbs_result: MultiChannelPRBSResult instance.
        output_path: Destination .xlsx file path.

    Returns:
        Path: Output file path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # Sheet 1: PRBS Samples
    # -------------------------------------------------------------
    ws_samples = wb.active
    ws_samples.title = "PRBS Samples"

    channels = [ch for ch in REQUIRED_CHANNELS if ch in prbs_result.channels]
    header = ["sample_index"]
    for ch in channels:
        header.extend(
            [
                f"{ch}_level",
                f"{ch}_seed_2bit",
                f"{ch}_prbs_2bit",
                f"{ch}_zero_state_2bit",
                f"{ch}_seed_4bit",
                f"{ch}_prbs_4bit",
                f"{ch}_zero_state_4bit",
            ]
        )
    ws_samples.append(header)

    for i in range(prbs_result.sample_count):
        row_vals: list[Any] = [i]
        for ch in channels:
            s_res = prbs_result.channels[ch].samples[i]
            row_vals.extend(
                [
                    s_res.level,
                    s_res.seed_2bit,
                    s_res.prbs_2bit,
                    s_res.is_zero_state_2bit,
                    s_res.seed_4bit,
                    s_res.prbs_4bit,
                    s_res.is_zero_state_4bit,
                ]
            )
        ws_samples.append(row_vals)
        row_idx = ws_samples.max_row

        # Format string columns explicitly to prevent Excel numeric conversion
        col_offset = 2
        for _ in channels:
            # col_offset is level (int)
            # col_offset+1 is seed_2bit (str)
            c_s2 = ws_samples.cell(row=row_idx, column=col_offset + 1)
            c_s2.data_type = "s"
            c_s2.number_format = "@"

            # col_offset+2 is prbs_2bit (str)
            c_p2 = ws_samples.cell(row=row_idx, column=col_offset + 2)
            c_p2.data_type = "s"
            c_p2.number_format = "@"

            # col_offset+4 is seed_4bit (str)
            c_s4 = ws_samples.cell(row=row_idx, column=col_offset + 4)
            c_s4.data_type = "s"
            c_s4.number_format = "@"

            # col_offset+5 is prbs_4bit (str)
            c_p4 = ws_samples.cell(row=row_idx, column=col_offset + 5)
            c_p4.data_type = "s"
            c_p4.number_format = "@"

            col_offset += 7

    # -------------------------------------------------------------
    # Sheet 2: PRBS Bitstreams (Chunked for Excel 32,767 cell limit)
    # -------------------------------------------------------------
    ws_bitstreams = wb.create_sheet(title="PRBS Bitstreams")
    ws_bitstreams.append(
        ["channel", "encoding_path", "bit_length", "part", "total_parts", "bitstream"]
    )

    for ch_name, ch_res in prbs_result.channels.items():
        for path_name, stream in [
            ("2-bit Gray PRBS (m=2)", ch_res.bitstream_2bit),
            ("4-bit PRBS (m=4)", ch_res.bitstream_4bit),
        ]:
            total_len = len(stream)
            if total_len <= EXCEL_MAX_CELL_CHARS:
                ws_bitstreams.append([ch_name, path_name, total_len, 1, 1, stream])
                c = ws_bitstreams.cell(row=ws_bitstreams.max_row, column=6)
                c.data_type = "s"
                c.number_format = "@"
            else:
                num_parts = (total_len + EXCEL_MAX_CELL_CHARS - 1) // EXCEL_MAX_CELL_CHARS
                for p_idx in range(num_parts):
                    chunk = stream[
                        p_idx * EXCEL_MAX_CELL_CHARS : (p_idx + 1) * EXCEL_MAX_CELL_CHARS
                    ]
                    ws_bitstreams.append(
                        [ch_name, path_name, total_len, p_idx + 1, num_parts, chunk]
                    )
                    c = ws_bitstreams.cell(row=ws_bitstreams.max_row, column=6)
                    c.data_type = "s"
                    c.number_format = "@"

    # -------------------------------------------------------------
    # Sheet 3: KAR Comparison
    # -------------------------------------------------------------
    ws_kar = wb.create_sheet(title="KAR Comparison")
    ws_kar.append(
        [
            "Pair",
            "Samples",
            "Level Agreement",
            "2-bit Pre-PRBS KAR",
            "2-bit Post-PRBS KAR",
            "2-bit KAR Delta",
            "2-bit Bits (Pre / Post)",
            "4-bit Pre-PRBS KAR",
            "4-bit Post-PRBS KAR",
            "4-bit KAR Delta",
            "4-bit Bits (Pre / Post)",
        ]
    )

    for pm in prbs_result.pair_metrics:
        ws_kar.append(
            [
                pm["pair"],
                pm["sample_count"],
                f"{float(pm['level_agreement_rate']):.4f}%",
                f"{float(pm['kar_2bit_before']):.4f}%",
                f"{float(pm['kar_2bit_after']):.4f}%",
                f"{float(pm['kar_2bit_delta']):+.4f}%",
                f"{pm['bit_count_2bit_before']} / {pm['bit_count_2bit_after']}",
                f"{float(pm['kar_4bit_before']):.4f}%",
                f"{float(pm['kar_4bit_after']):.4f}%",
                f"{float(pm['kar_4bit_delta']):+.4f}%",
                f"{pm['bit_count_4bit_before']} / {pm['bit_count_4bit_after']}",
            ]
        )

    # -------------------------------------------------------------
    # Sheet 4: Configuration & Zero Seeds
    # -------------------------------------------------------------
    ws_cfg = wb.create_sheet(title="Configuration & Zero Seeds")
    ws_cfg.append(["Parameter / Metric", "Value", "Description"])
    ws_cfg.append(["Algorithm", "Per-Sample Galois XOR-LFSR PRBS Expansion", "PA.pdf Section 3.2.4"])
    ws_cfg.append(["Shift Direction", prbs_result.config.shift_direction, "MSB emitted first"])
    ws_cfg.append(["Bit Ordering", "Big-endian", "Index 0 of string seed is MSB"])
    ws_cfg.append(["2-bit Polynomial", prbs_result.config.poly_2bit_name, "Primitive polynomial over GF(2)"])
    ws_cfg.append(["2-bit Galois Mask", f"0b{prbs_result.config.mask_2bit:02b} ({prbs_result.config.mask_2bit})", "Taps x^1 + x^0"])
    ws_cfg.append(["2-bit Bits/Sample", prbs_result.config.output_length_2bit, "2^2 - 1 = 3 bits"])
    ws_cfg.append(["4-bit Polynomial", prbs_result.config.poly_4bit_name, "Primitive polynomial over GF(2)"])
    ws_cfg.append(["4-bit Galois Mask", f"0b{prbs_result.config.mask_4bit:04b} ({prbs_result.config.mask_4bit})", "Taps x^1 + x^0"])
    ws_cfg.append(["4-bit Bits/Sample", prbs_result.config.output_length_4bit, "2^4 - 1 = 15 bits"])
    ws_cfg.append(["Zero Seed Handling", "Standard evolution -> 000 (period 1)", "NEVER remapped; deliberate exception to PA.pdf"])
    ws_cfg.append(["Total Samples", prbs_result.sample_count, "Number of observations"])

    ws_cfg.append([])
    ws_cfg.append(["Channel", "2-bit Zero Seeds (Seed '00')", "4-bit Zero Seeds (Seed '0000')"])
    for ch_name, ch_res in prbs_result.channels.items():
        ws_cfg.append([ch_name, ch_res.zero_seed_count_2bit, ch_res.zero_seed_count_4bit])

    wb.save(out_path)
    return out_path
