"""Modified Adaptive Dual-Threshold Quantization (MADQ) module.

Governing Specification:
    PA.pdf Section 3.2.3, Equations (3.21)–(3.23), pages 27–28.
    PA.pdf Section 3.2.8.c, Equation (3.64), page 35 (Key Agreement Rate).

Mathematical Formulation:
    For each channel and non-overlapping segment of length W (e.g., W = 50):
        mean = (1 / M) * sum(X_i)
        std = sqrt((1 / M) * sum((X_i - mean)^2))  (population std, ddof = 0)
        upper threshold q_i^+ = mean + alpha * std  (Equation 3.21)
        lower threshold q_i^- = mean - alpha * std  (Equation 3.22)

    Quantization level assignment (Equation 3.23):
        Q(X_i) = 0, if X_i < q_i^-
        Q(X_i) = 1, if q_i^- <= X_i <= q_i^+
        Q(X_i) = 2, if X_i > q_i^+

    Every sample is retained. Constant segments (std = 0) produce level 1.

Binary Encodings:
    - 2-bit Gray code (implementation choice):
        0 -> "00"
        1 -> "01"
        2 -> "11"
    - 4-bit code (PA.pdf Section 3.2.3):
        0 -> "1010"
        1 -> "1011"
        2 -> "1001"

Evaluation Metric:
    Bit-level Key Agreement Rate (KAR) per Equation (3.64):
        KAR = ((N_total - N_mismatch) / N_total) * 100%
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
import json
import math
from pathlib import Path
from typing import Any, Mapping, Sequence

GRAY_2BIT_MAP: dict[int, str] = {0: "00", 1: "01", 2: "11"}
CODE_4BIT_MAP: dict[int, str] = {0: "1010", 1: "1011", 2: "1001"}

REQUIRED_CHANNELS: tuple[str, ...] = ("Alice", "Bob", "Eve1-Alice", "Eve1-Bob")
EVALUATION_PAIRS: tuple[tuple[str, str, str], ...] = (
    ("Alice", "Bob", "Alice–Bob"),
    ("Alice", "Eve1-Alice", "Alice–Eve1-Alice"),
    ("Bob", "Eve1-Bob", "Bob–Eve1-Bob"),
)


@dataclass(frozen=True)
class QuantizationConfig:
    """Configuration parameters for Modified Adaptive Dual-Threshold Quantization.

    Attributes:
        window_size: Number of observations per non-overlapping segment.
            Provisional software-validation setting: 50. Must be a positive integer.
        alpha: Threshold multiplier coefficient alpha.
            Provisional software-validation setting: 0.5. Must be a finite non-negative float.
    """

    window_size: int = 50
    alpha: float = 0.5

    def __post_init__(self) -> None:
        if isinstance(self.window_size, bool) or not isinstance(self.window_size, int):
            raise TypeError(
                f"window_size must be an integer, got {type(self.window_size).__name__}."
            )
        if self.window_size <= 0:
            raise ValueError(
                f"window_size must be a positive integer (> 0), got {self.window_size}."
            )

        if isinstance(self.alpha, bool) or not isinstance(self.alpha, (int, float)):
            raise TypeError(
                f"alpha must be a real number, got {type(self.alpha).__name__}."
            )
        if not math.isfinite(self.alpha):
            raise ValueError(f"alpha must be finite, got {self.alpha}.")
        if self.alpha < 0.0:
            raise ValueError(f"alpha must be non-negative (>= 0), got {self.alpha}.")

    def to_metadata_dict(self) -> dict[str, Any]:
        """Convert configuration to structured metadata dictionary."""
        return {
            "algorithm": "Modified Adaptive Dual-Threshold Quantization (MADQ)",
            "governing_specification": {
                "thresholds_and_levels": "PA.pdf Section 3.2.3, Equations (3.21)–(3.23), pages 27–28",
                "key_agreement_rate": "PA.pdf Section 3.2.8.c, Equation (3.64), page 35",
            },
            "parameters": {
                "window_size": self.window_size,
                "alpha": float(self.alpha),
                "ddof": 0,
                "segmentation": "non-overlapping",
                "status": "provisional software-validation settings (not optimized physical-channel parameters)",
            },
            "encodings": {
                "gray_2bit": {
                    "mapping": {"0": "00", "1": "01", "2": "11"},
                    "description": "2-bit Gray code representation (implementation choice)",
                },
                "code_4bit": {
                    "mapping": {"0": "1010", "1": "1011", "2": "1001"},
                    "description": "4-bit representation (PA.pdf Section 3.2.3)",
                },
            },
        }


@dataclass(frozen=True)
class SegmentThresholds:
    """Thresholds and statistics for a single segment of a channel."""

    segment_index: int
    channel: str
    start_index: int
    end_index: int
    sample_count: int
    mean: float
    std: float
    lower_threshold: float
    upper_threshold: float


@dataclass
class ChannelQuantizationResult:
    """Container for multi-channel quantization results."""

    config: QuantizationConfig
    sample_count: int
    num_segments: int
    levels: dict[str, list[int]]
    thresholds: dict[str, list[SegmentThresholds]]
    bitstreams_2bit: dict[str, str]
    bitstreams_4bit: dict[str, str]
    level_frequencies: dict[str, dict[int, dict[str, float | int]]]
    pair_metrics: list[dict[str, Any]]
    figure_paths: tuple[Path, Path] | None = None


def quantize_series(
    series: Sequence[float | int],
    config: QuantizationConfig | None = None,
    channel_name: str = "channel",
) -> tuple[list[int], list[SegmentThresholds]]:
    """Quantize a 1D sequence of observations using Modified ADQ.

    Args:
        series: 1D sequence of numeric observations.
        config: QuantizationConfig instance. Uses defaults if None.
        channel_name: Identifier for the channel being quantized.

    Returns:
        tuple: (levels, segment_thresholds) where levels is a list of int (0, 1, 2)
            and segment_thresholds is a list of SegmentThresholds.

    Raises:
        ValueError: If series is empty, contains non-numeric or non-finite values.
    """
    cfg = config if config is not None else QuantizationConfig()
    n_samples = len(series)

    if n_samples == 0:
        raise ValueError("Input series must not be empty.")

    # Validate finite numeric values
    float_series: list[float] = []
    for idx, val in enumerate(series):
        try:
            f_val = float(val)
        except (TypeError, ValueError) as err:
            raise ValueError(
                f"Invalid non-numeric value in channel '{channel_name}' at index {idx}: {val!r}."
            ) from err

        if not math.isfinite(f_val):
            raise ValueError(
                f"Non-finite value (NaN/Inf) encountered in channel '{channel_name}' at index {idx}: {f_val}."
            )
        float_series.append(f_val)

    window_size = cfg.window_size
    alpha = cfg.alpha

    num_segments = (n_samples + window_size - 1) // window_size
    levels: list[int] = []
    segments: list[SegmentThresholds] = []

    for seg_idx in range(num_segments):
        start = seg_idx * window_size
        end = min(start + window_size, n_samples)
        seg_vals = float_series[start:end]
        seg_count = len(seg_vals)

        seg_mean = math.fsum(seg_vals) / float(seg_count)
        # Population standard deviation (ddof=0)
        variance = math.fsum((x - seg_mean) ** 2 for x in seg_vals) / float(seg_count)
        seg_std = math.sqrt(max(0.0, variance))

        lower = seg_mean - alpha * seg_std
        upper = seg_mean + alpha * seg_std

        is_constant = (seg_std == 0.0)

        seg_threshold = SegmentThresholds(
            segment_index=seg_idx,
            channel=channel_name,
            start_index=start,
            end_index=end,
            sample_count=seg_count,
            mean=seg_mean,
            std=seg_std,
            lower_threshold=lower,
            upper_threshold=upper,
        )
        segments.append(seg_threshold)

        for x in seg_vals:
            if is_constant:
                levels.append(1)
            elif x < lower:
                levels.append(0)
            elif x > upper:
                levels.append(2)
            else:
                levels.append(1)

    return levels, segments


def encode_levels_to_bits(
    levels: Sequence[int],
    mapping: Mapping[int, str],
) -> str:
    """Encode quantization levels (0, 1, 2) into a continuous binary bitstream.

    Sample order is strictly preserved.

    Args:
        levels: Sequence of integer quantization levels.
        mapping: Dictionary mapping level int to binary string.

    Returns:
        str: Continuous binary string.

    Raises:
        ValueError: If an unexpected level value is encountered.
    """
    bit_chunks: list[str] = []
    for idx, lvl in enumerate(levels):
        if lvl not in mapping:
            raise ValueError(
                f"Invalid quantization level {lvl!r} at sample index {idx}; expected one of {list(mapping.keys())}."
            )
        bit_chunks.append(mapping[lvl])
    return "".join(bit_chunks)


def calculate_kar(bitstream_a: str, bitstream_b: str) -> float:
    """Calculate Key Agreement Rate (KAR) between two bitstreams per PA.pdf Eq (3.64).

    Equation (3.64):
        KAR = ((N_total - N_mismatch) / N_total) * 100%

    Args:
        bitstream_a: First binary string.
        bitstream_b: Second binary string.

    Returns:
        float: Key Agreement Rate percentage in [0.0, 100.0].

    Raises:
        ValueError: If bitstreams have unequal length, are empty, or contain non-binary characters.
    """
    len_a = len(bitstream_a)
    len_b = len(bitstream_b)

    if len_a != len_b:
        raise ValueError(
            f"Bitstreams must have identical length: len(a)={len_a}, len(b)={len_b}."
        )
    if len_a == 0:
        raise ValueError("Bitstreams must not be empty.")

    n_match = 0
    for idx in range(len_a):
        bit_a = bitstream_a[idx]
        bit_b = bitstream_b[idx]
        if bit_a not in ("0", "1") or bit_b not in ("0", "1"):
            raise ValueError(
                f"Non-binary character encountered at bit index {idx}: a={bit_a!r}, b={bit_b!r}."
            )
        if bit_a == bit_b:
            n_match += 1

    return (float(n_match) / float(len_a)) * 100.0


def calculate_level_agreement(
    levels_a: Sequence[int],
    levels_b: Sequence[int],
) -> float:
    """Calculate sample-level agreement rate between two quantization level sequences.

    Args:
        levels_a: First sequence of quantization levels.
        levels_b: Second sequence of quantization levels.

    Returns:
        float: Sample-level agreement percentage in [0.0, 100.0].

    Raises:
        ValueError: If sequences have unequal length or are empty.
    """
    n_a = len(levels_a)
    n_b = len(levels_b)

    if n_a != n_b:
        raise ValueError(
            f"Level sequences must have identical length: len(a)={n_a}, len(b)={n_b}."
        )
    if n_a == 0:
        raise ValueError("Level sequences must not be empty.")

    matches = sum(1 for a, b in zip(levels_a, levels_b) if a == b)
    return (float(matches) / float(n_a)) * 100.0


def calculate_level_frequencies(
    levels: Sequence[int],
) -> dict[int, dict[str, float | int]]:
    """Compute count and percentage for quantization levels 0, 1, and 2.

    Args:
        levels: Sequence of integer quantization levels.

    Returns:
        dict: Mapping of level (0, 1, 2) to {"count": int, "percentage": float}.
    """
    n = len(levels)
    counts = {0: 0, 1: 0, 2: 0}
    for idx, lvl in enumerate(levels):
        if lvl not in counts:
            raise ValueError(
                f"Invalid quantization level {lvl!r} at index {idx}; expected 0, 1, or 2."
            )
        counts[lvl] += 1

    return {
        lvl: {
            "count": counts[lvl],
            "percentage": (float(counts[lvl]) / float(n) * 100.0) if n > 0 else 0.0,
        }
        for lvl in (0, 1, 2)
    }


def quantize_channels(
    channels: Mapping[str, Sequence[float | int]],
    config: QuantizationConfig | None = None,
) -> ChannelQuantizationResult:
    """Quantize multiple RSSI channels independently using Modified ADQ.

    Strict channel isolation is enforced: each channel's mean and standard deviation
    are computed locally within each segment without cross-channel leakage.

    Args:
        channels: Mapping of channel name to observation series.
        config: QuantizationConfig instance (uses defaults if None).

    Returns:
        ChannelQuantizationResult: Fully populated quantization result container.

    Raises:
        ValueError: If channels mapping is empty, channels have unequal lengths,
            or observations are invalid.
    """
    cfg = config if config is not None else QuantizationConfig()

    if not channels:
        raise ValueError("Channels mapping must not be empty.")

    channel_names = list(channels.keys())
    first_ch = channel_names[0]
    expected_length = len(channels[first_ch])

    if expected_length == 0:
        raise ValueError(f"Channel '{first_ch}' has zero samples.")

    for ch in channel_names[1:]:
        ch_len = len(channels[ch])
        if ch_len != expected_length:
            raise ValueError(
                f"Channel length mismatch: '{first_ch}' has {expected_length} samples, "
                f"but '{ch}' has {ch_len} samples."
            )

    levels_dict: dict[str, list[int]] = {}
    thresholds_dict: dict[str, list[SegmentThresholds]] = {}
    bitstreams_2bit: dict[str, str] = {}
    bitstreams_4bit: dict[str, str] = {}
    level_freqs: dict[str, dict[int, dict[str, float | int]]] = {}

    num_segments = (expected_length + cfg.window_size - 1) // cfg.window_size

    for ch in channel_names:
        ch_levels, ch_thresholds = quantize_series(
            series=channels[ch],
            config=cfg,
            channel_name=ch,
        )
        levels_dict[ch] = ch_levels
        thresholds_dict[ch] = ch_thresholds

        bitstreams_2bit[ch] = encode_levels_to_bits(ch_levels, GRAY_2BIT_MAP)
        bitstreams_4bit[ch] = encode_levels_to_bits(ch_levels, CODE_4BIT_MAP)
        level_freqs[ch] = calculate_level_frequencies(ch_levels)

    # Evaluate pair metrics
    pair_metrics: list[dict[str, Any]] = []
    for col1, col2, pair_label in EVALUATION_PAIRS:
        if col1 in levels_dict and col2 in levels_dict:
            lvl_agr = calculate_level_agreement(levels_dict[col1], levels_dict[col2])
            kar_2bit = calculate_kar(bitstreams_2bit[col1], bitstreams_2bit[col2])
            kar_4bit = calculate_kar(bitstreams_4bit[col1], bitstreams_4bit[col2])

            pair_metrics.append(
                {
                    "pair": pair_label,
                    "channel_a": col1,
                    "channel_b": col2,
                    "sample_count": expected_length,
                    "level_agreement_rate": lvl_agr,
                    "kar_2bit": kar_2bit,
                    "kar_4bit": kar_4bit,
                    "bit_count_2bit": len(bitstreams_2bit[col1]),
                    "bit_count_4bit": len(bitstreams_4bit[col1]),
                }
            )

    return ChannelQuantizationResult(
        config=cfg,
        sample_count=expected_length,
        num_segments=num_segments,
        levels=levels_dict,
        thresholds=thresholds_dict,
        bitstreams_2bit=bitstreams_2bit,
        bitstreams_4bit=bitstreams_4bit,
        level_frequencies=level_freqs,
        pair_metrics=pair_metrics,
    )


def export_quantized_samples_csv(
    filtered_data: Mapping[str, Sequence[float]],
    quant_result: ChannelQuantizationResult,
    output_path: str | Path,
) -> Path:
    """Export aligned sample-level quantization results to CSV.

    Columns include for each channel: filtered RSSI, segment index, lower threshold,
    upper threshold, quantization level, 2-bit Gray encoding, and 4-bit encoding,
    along with sample-level agreement flags.

    Args:
        filtered_data: Mapping of channel name to filtered RSSI observations.
        quant_result: ChannelQuantizationResult instance.
        output_path: Path to destination CSV.

    Returns:
        Path: Resolved destination path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    columns = [ch for ch in REQUIRED_CHANNELS if ch in quant_result.levels]
    n_samples = quant_result.sample_count
    w = quant_result.config.window_size

    header = ["sample_index", "segment_index"]
    for col in columns:
        header.extend([
            f"{col}_filtered",
            f"{col}_lower_threshold",
            f"{col}_upper_threshold",
            f"{col}_level",
            f"{col}_bits_2bit",
            f"{col}_bits_4bit",
        ])
    if "Alice" in quant_result.levels and "Bob" in quant_result.levels:
        header.append("alice_bob_level_match")

    # Map segment thresholds by (channel, seg_idx) for fast lookup
    thresh_map: dict[tuple[str, int], SegmentThresholds] = {}
    for ch, seg_list in quant_result.thresholds.items():
        for st in seg_list:
            thresh_map[(ch, st.segment_index)] = st

    import io

    sio = io.StringIO()
    writer = csv.writer(sio)
    writer.writerow(header)

    for i in range(n_samples):
        seg_idx = i // w
        row: list[Any] = [i, seg_idx]
        for col in columns:
            st = thresh_map[(col, seg_idx)]
            lvl = quant_result.levels[col][i]
            b2 = GRAY_2BIT_MAP[lvl]
            b4 = CODE_4BIT_MAP[lvl]
            filt_val = filtered_data[col][i]
            row.extend([
                filt_val,
                st.lower_threshold,
                st.upper_threshold,
                lvl,
                b2,
                b4,
            ])
        if "Alice" in quant_result.levels and "Bob" in quant_result.levels:
            match_flag = 1 if quant_result.levels["Alice"][i] == quant_result.levels["Bob"][i] else 0
            row.append(match_flag)
        writer.writerow(row)

    content = sio.getvalue()
    try:
        with open(out_path, mode="w", newline="", encoding="utf-8") as f:
            f.write(content)
    except PermissionError as p_err:
        if out_path.is_file():
            try:
                existing = out_path.read_text(encoding="utf-8")
                if existing.replace("\r\n", "\n") == content.replace("\r\n", "\n"):
                    return out_path
            except Exception:
                pass
        raise PermissionError(
            f"Cannot write to '{out_path}'. The file appears to be open and locked by another application (e.g. Excel). "
            f"Please close the file in Excel and retry."
        ) from p_err

    return out_path



def export_quantized_segments_csv(
    quant_result: ChannelQuantizationResult,
    output_path: str | Path,
) -> Path:
    """Export segment thresholds and statistics to CSV.

    Args:
        quant_result: ChannelQuantizationResult instance.
        output_path: Path to destination CSV.

    Returns:
        Path: Resolved destination path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    header = [
        "segment_index",
        "channel",
        "start_index",
        "end_index",
        "sample_count",
        "mean",
        "std",
        "lower_threshold",
        "upper_threshold",
    ]

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)

        for seg_idx in range(quant_result.num_segments):
            for ch in REQUIRED_CHANNELS:
                if ch in quant_result.thresholds:
                    st = quant_result.thresholds[ch][seg_idx]
                    writer.writerow([
                        st.segment_index,
                        st.channel,
                        st.start_index,
                        st.end_index,
                        st.sample_count,
                        st.mean,
                        st.std,
                        st.lower_threshold,
                        st.upper_threshold,
                    ])

    return out_path


def export_quantized_bitstreams_csv(
    quant_result: ChannelQuantizationResult,
    output_path: str | Path,
) -> Path:
    """Export encoded bitstreams to CSV.

    Args:
        quant_result: ChannelQuantizationResult instance.
        output_path: Path to destination CSV.

    Returns:
        Path: Resolved destination path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    header = ["channel", "encoding", "bit_count", "bitstream"]

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)

        for ch in REQUIRED_CHANNELS:
            if ch in quant_result.bitstreams_2bit:
                bs2 = quant_result.bitstreams_2bit[ch]
                writer.writerow([ch, "2-bit Gray", len(bs2), bs2])
            if ch in quant_result.bitstreams_4bit:
                bs4 = quant_result.bitstreams_4bit[ch]
                writer.writerow([ch, "4-bit", len(bs4), bs4])

    return out_path


def export_quantization_metrics_csv(
    quant_result: ChannelQuantizationResult,
    output_path: str | Path,
) -> Path:
    """Export Key Agreement Rate (KAR) and Level Agreement metrics to CSV.

    Args:
        quant_result: ChannelQuantizationResult instance.
        output_path: Path to destination CSV.

    Returns:
        Path: Resolved destination path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = [
        "pair",
        "sample_count",
        "level_agreement_rate",
        "kar_2bit",
        "kar_4bit",
        "bit_count_2bit",
        "bit_count_4bit",
    ]

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in quant_result.pair_metrics:
            csv_row = {k: row[k] for k in fieldnames}
            writer.writerow(csv_row)

    return out_path


def export_quantization_metadata_json(
    quant_result: ChannelQuantizationResult,
    output_path: str | Path,
    extra_metadata: dict[str, Any] | None = None,
) -> Path:
    """Export complete quantization configuration, rationale, and metrics to JSON.

    Args:
        quant_result: ChannelQuantizationResult instance.
        output_path: Path to destination JSON.
        extra_metadata: Optional dictionary of extra metadata fields.

    Returns:
        Path: Resolved destination path.
    """
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    meta = quant_result.config.to_metadata_dict()
    meta["execution_summary"] = {
        "sample_count": quant_result.sample_count,
        "num_segments": quant_result.num_segments,
        "pair_metrics": quant_result.pair_metrics,
        "level_frequencies": quant_result.level_frequencies,
    }
    meta["bitstreams"] = {
        "2bit_gray": quant_result.bitstreams_2bit,
        "4bit": quant_result.bitstreams_4bit,
    }

    if extra_metadata:
        meta.update(extra_metadata)

    with open(out_path, mode="w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    return out_path


EXCEL_MAX_CELL_CHARS: int = 32767


def export_quantized_workbook_xlsx(
    filtered_data: Mapping[str, Sequence[float]],
    quant_result: ChannelQuantizationResult,
    output_path: str | Path,
    max_cell_chars: int = EXCEL_MAX_CELL_CHARS,
) -> Path:
    """Export complete quantization results to an Excel workbook (.xlsx).

    Preserves exact binary strings (e.g. '00', '01') by explicitly storing them
    as text-formatted string cells ('@' number format). This avoids Excel's
    numeric auto-conversion bug where leading zeros in CSVs are dropped
    ('00' -> 0, '01' -> 1).

    Sheets included:
        1. 'Quantized Samples': Aligned sample-level filtered values, thresholds,
           levels, 2-bit Gray codes, 4-bit codes, and Alice-Bob match flag.
        2. 'Bitstreams': Channel bitstreams for 2-bit Gray and 4-bit encodings.
           If a bitstream exceeds Excel's 32,767-character limit, it is split
           into sequential chunks across multiple rows with part indicators to
           prevent silent truncation and preserve 100% of the key bits.
        3. 'Agreement Metrics': Pairwise Key Agreement Rate (KAR), Level Agreement,
           and bit counts.
        4. 'Segment Thresholds': Statistical moments and thresholds per segment.

    Args:
        filtered_data: Mapping of channel name to filtered RSSI observations.
        quant_result: ChannelQuantizationResult instance.
        output_path: Path to destination .xlsx workbook.
        max_cell_chars: Maximum character capacity per cell (default: 32,767).

    Returns:
        Path: Resolved destination path.

    Raises:
        ValueError: If max_cell_chars is non-positive or if data cannot be preserved.
    """
    import openpyxl
    from openpyxl.styles import Font

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if max_cell_chars <= 0:
        raise ValueError(f"max_cell_chars must be a positive integer, got {max_cell_chars}.")

    wb = openpyxl.Workbook()
    header_font = Font(bold=True)

    # -------------------------------------------------------------
    # Sheet 1: Quantized Samples
    # -------------------------------------------------------------
    ws_samples = wb.active
    ws_samples.title = "Quantized Samples"

    columns = [ch for ch in REQUIRED_CHANNELS if ch in quant_result.levels]
    n_samples = quant_result.sample_count
    w = quant_result.config.window_size

    samples_header = ["sample_index", "segment_index"]
    for col in columns:
        samples_header.extend([
            f"{col}_filtered",
            f"{col}_lower_threshold",
            f"{col}_upper_threshold",
            f"{col}_level",
            f"{col}_bits_2bit",
            f"{col}_bits_4bit",
        ])
    if "Alice" in quant_result.levels and "Bob" in quant_result.levels:
        samples_header.append("alice_bob_level_match")

    ws_samples.append(samples_header)
    for cell in ws_samples[1]:
        cell.font = header_font

    thresh_map: dict[tuple[str, int], SegmentThresholds] = {}
    for ch, seg_list in quant_result.thresholds.items():
        for st in seg_list:
            thresh_map[(ch, st.segment_index)] = st

    for i in range(n_samples):
        seg_idx = i // w
        row: list[Any] = [i, seg_idx]
        for col in columns:
            st = thresh_map[(col, seg_idx)]
            lvl = quant_result.levels[col][i]
            b2 = GRAY_2BIT_MAP[lvl]
            b4 = CODE_4BIT_MAP[lvl]
            filt_val = filtered_data[col][i]
            row.extend([
                filt_val,
                st.lower_threshold,
                st.upper_threshold,
                lvl,
                b2,
                b4,
            ])
        if "Alice" in quant_result.levels and "Bob" in quant_result.levels:
            match_flag = 1 if quant_result.levels["Alice"][i] == quant_result.levels["Bob"][i] else 0
            row.append(match_flag)

        ws_samples.append(row)
        curr_row = ws_samples.max_row
        for k in range(len(columns)):
            c_b2 = ws_samples.cell(row=curr_row, column=7 + k * 6)
            c_b2.data_type = "s"
            c_b2.number_format = "@"

            c_b4 = ws_samples.cell(row=curr_row, column=8 + k * 6)
            c_b4.data_type = "s"
            c_b4.number_format = "@"

    # -------------------------------------------------------------
    # Sheet 2: Bitstreams
    # -------------------------------------------------------------
    ws_bits = wb.create_sheet(title="Bitstreams")
    bits_header = ["channel", "encoding", "part", "total_parts", "bit_count", "bitstream"]
    ws_bits.append(bits_header)
    for cell in ws_bits[1]:
        cell.font = header_font

    for ch in REQUIRED_CHANNELS:
        enc_entries = []
        if ch in quant_result.bitstreams_2bit:
            enc_entries.append(("2-bit Gray", quant_result.bitstreams_2bit[ch]))
        if ch in quant_result.bitstreams_4bit:
            enc_entries.append(("4-bit", quant_result.bitstreams_4bit[ch]))

        for enc_name, bs in enc_entries:
            bs_len = len(bs)
            if bs_len <= max_cell_chars:
                ws_bits.append([ch, enc_name, 1, 1, bs_len, bs])
                c = ws_bits.cell(row=ws_bits.max_row, column=6)
                c.data_type = "s"
                c.number_format = "@"
            else:
                total_parts = math.ceil(bs_len / max_cell_chars)
                for p_idx in range(total_parts):
                    chunk = bs[p_idx * max_cell_chars : (p_idx + 1) * max_cell_chars]
                    if len(chunk) > max_cell_chars:
                        raise ValueError(
                            f"Bitstream chunk exceeds maximum cell limit of {max_cell_chars} characters."
                        )
                    ws_bits.append([ch, enc_name, p_idx + 1, total_parts, len(chunk), chunk])
                    c = ws_bits.cell(row=ws_bits.max_row, column=6)
                    c.data_type = "s"
                    c.number_format = "@"

    # -------------------------------------------------------------
    # Sheet 3: Agreement Metrics
    # -------------------------------------------------------------
    ws_metrics = wb.create_sheet(title="Agreement Metrics")
    metrics_header = [
        "pair",
        "channel_a",
        "channel_b",
        "sample_count",
        "level_agreement_rate",
        "kar_2bit",
        "kar_4bit",
        "bit_count_2bit",
        "bit_count_4bit",
    ]
    ws_metrics.append(metrics_header)
    for cell in ws_metrics[1]:
        cell.font = header_font

    for pm in quant_result.pair_metrics:
        ws_metrics.append([
            pm["pair"],
            pm["channel_a"],
            pm["channel_b"],
            pm["sample_count"],
            pm["level_agreement_rate"],
            pm["kar_2bit"],
            pm["kar_4bit"],
            pm["bit_count_2bit"],
            pm["bit_count_4bit"],
        ])

    # -------------------------------------------------------------
    # Sheet 4: Segment Thresholds
    # -------------------------------------------------------------
    ws_segs = wb.create_sheet(title="Segment Thresholds")
    segs_header = [
        "segment_index",
        "channel",
        "start_index",
        "end_index",
        "sample_count",
        "mean",
        "std",
        "lower_threshold",
        "upper_threshold",
    ]
    ws_segs.append(segs_header)
    for cell in ws_segs[1]:
        cell.font = header_font

    for seg_idx in range(quant_result.num_segments):
        for ch in REQUIRED_CHANNELS:
            if ch in quant_result.thresholds:
                st = quant_result.thresholds[ch][seg_idx]
                ws_segs.append([
                    st.segment_index,
                    st.channel,
                    st.start_index,
                    st.end_index,
                    st.sample_count,
                    st.mean,
                    st.std,
                    st.lower_threshold,
                    st.upper_threshold,
                ])

    wb.save(out_path)
    return out_path
